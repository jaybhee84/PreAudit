from pathlib import Path
from collections import defaultdict
import re, unicodedata, math
import openpyxl

DEFAULT_FOLDER = r'C:\Users\DepEd\Downloads\PBB 2024 WITH IUs'

def norm(v):
    s = ''.join(c for c in unicodedata.normalize('NFKD', str(v or '')) if not unicodedata.combining(c)).upper()
    return ' '.join(re.findall(r'[A-Z0-9]+', s))

def number(v):
    if isinstance(v, bool): return None
    if isinstance(v, (int, float)) and math.isfinite(v): return float(v)
    return None

def integer(v):
    if number(v) is not None: return int(v) if float(v).is_integer() else None
    m = re.fullmatch(r'\s*(?:SG|step)?\s*(\d+)\s*,?\s*', str(v), re.I)
    return int(m[1]) if m else None

def shortkey(v):
    words = norm(v).split()
    return (words[0], words[1]) if ',' in str(v) and len(words)>1 else None

def audit(folder):
    root = Path(folder)
    if not root.is_dir(): raise ValueError('Choose an existing folder.')
    paths = sorted(p for p in root.rglob('*') if p.suffix.lower() in ('.xlsx','.xlsm') and not p.name.startswith('~$'))
    if not paths: raise ValueError('No Excel workbooks found.')
    issues=[]; files=[]; records=[]; personnel=defaultdict(list); separations=[]
    def issue(level, rule, file, sheet='', cell='', employee='', detail=''):
        issues.append(dict(level=level,rule=rule,file=file,sheet=sheet,cell=cell,employee=employee,detail=detail))
    for path in paths:
        file=str(path.relative_to(root))
        try:
            w=openpyxl.load_workbook(path,read_only=True,data_only=True)
            try: sheets={s.title:list(s.values) for s in w}
            finally: w.close()
        except Exception as e:
            issue('Error','Unreadable workbook',file,detail=str(e)); files.append(dict(file=file,role='Unreadable',records=0)); continue
        if 'LIST OF PERSONNEL' in path.name.upper():
            for sn,rows in sheets.items():
                for r,row in enumerate(rows,1):
                    if r>2 and len(row)>7 and row[4] and row[2]: personnel[norm(row[4])].append((file,sn,r,row))
            files.append(dict(file=file,role='Personnel reference',records=sum(len(v) for v in personnel.values())))
            continue
        if 'SEPARATION' in path.name.upper():
            n=0
            for sn,rows in sheets.items():
                if 'DIRECTORY' not in sn.upper(): continue
                for r,row in enumerate(rows,1):
                    if r>12 and len(row)>6 and row[1] and number(row[0]) is not None:
                        separations.append((file,sn,r,row)); n+=1
            files.append(dict(file=file,role='Separation reference',records=n)); continue
        count=0
        for sn,rows in sheets.items():
            header=next((i for i,row in enumerate(rows[:40]) if any('NAME OF EMPLOYEE' in str(v).upper() for v in row)),None)
            if header is None: continue
            cols={}
            for c,v in enumerate(rows[header]):
                t=str(v).upper()
                for key,phrase in [('name','NAME OF EMPLOYEE'),('sg','SG\n'),('step','STEP INCREMENT'),('months','MONTHS IN SERVICE'),('salary','MONTHLY BASIC SALARY'),('pbb','AMOUNT OF THE PBB')]:
                    if phrase in t: cols[key]=c
            if len(cols)!=6:
                issue('Error','Unsupported form layout',file,sn,detail='Required employee columns could not all be detected.');continue
            get=lambda row,key: row[cols[key]] if cols[key]<len(row) else None
            local=[]; unit=''
            summary_col=cols['sg']
            rate=number(rows[13][summary_col]) if len(rows)>13 and len(rows[13])>summary_col else None
            summary_letter=openpyxl.utils.get_column_letter(summary_col+1)
            rating={str(row[0]).strip():number(row[1]) for row in sheets.get('Rating',[]) if len(row)>1 and row[0] is not None and number(row[1]) is not None}
            salaryrows=sheets.get('2024_Tranche',[])
            for r,row in enumerate(rows[header+1:],header+2):
                text=' '.join(str(v) for v in row[:5] if v)
                if re.search(r'ineligible|not eligible',text,re.I): break
                name=get(row,'name')
                if len(row)>1 and row[1]: unit=str(row[1])
                sg=integer(get(row,'sg')); step=integer(get(row,'step'))
                if not name and (sg is not None or step is not None):
                    issue('Error','Missing employee name',file,sn,f'{openpyxl.utils.get_column_letter(cols["name"]+1)}{r}');continue
                if not name or (sg is None and step is None and not get(row,'months')): continue
                name=str(name); count+=1
                rec=dict(file=file,sheet=sn,row=r,name=name,unit=unit,sg=sg,step=step,salary=get(row,'salary'),pbb=get(row,'pbb'))
                local.append(rec);records.append(rec)
                def flag(rule,key,detail,level='Error'):
                    issue(level,rule,file,sn,f'{openpyxl.utils.get_column_letter(cols[key]+1)}{r}',name,detail)
                if not sg or not 1<=sg<=33: flag('Invalid salary grade','sg',f'Found {get(row,"sg")!s}; expected 1–33.')
                if not step or not 1<=step<=8: flag('Invalid step','step',f'Found {get(row,"step")!s}; expected 1–8.')
                factor=rating.get(str(get(row,'months')).strip())
                if factor is None: flag('Unrecognized service period','months',f'No matching factor in Rating table: {get(row,"months")!s}.','Review')
                salary=number(rec['salary']); pbb=number(rec['pbb'])
                if salary is None or salary<=0: flag('Missing or invalid salary','salary','Salary must be a positive number. Formula caches may need refreshing in Excel.')
                if pbb is None or pbb<0: flag('Missing or invalid PBB','pbb','PBB must be numeric and nonnegative. Formula caches may need refreshing in Excel.')
                if sg and step and 1<=sg<=33 and 1<=step<=8 and len(salaryrows)>sg and len(salaryrows[sg])>step:
                    expected=number(salaryrows[sg][step])
                    if expected is None: flag('Invalid salary reference','salary','The matching 2024_Tranche value is not numeric.','Review')
                    elif salary is not None and abs(salary-expected)>.01: flag('Salary table mismatch','salary',f'Workbook 2024_Tranche: {expected:,.2f}; form: {salary:,.2f}.')
                if salary is not None and pbb is not None and rate is not None and factor is not None:
                    expected=salary*rate*factor
                    if abs(pbb-expected)>.02: flag('PBB calculation mismatch','pbb',f'Salary × workbook rate ({rate:g}) × service factor ({factor:g}) = {expected:,.2f}; form: {pbb:,.2f}.')
                elif rate is None: flag('Missing agency rate','pbb',f'No numeric agency rate at {summary_letter}14.','Review')
            total=sum(number(x['pbb']) or 0 for x in local)
            complete=all(number(x['pbb']) is not None for x in local)
            for rr,expected,label in [(13,len(local),'Employee count'),(15,total,'PBB total')]:
                reported=number(rows[rr-1][summary_col]) if len(rows)>=rr and len(rows[rr-1])>summary_col else None
                if rr==15 and not complete:
                    issue('Review','Total unavailable',file,sn,f'{summary_letter}15',detail='Some employee PBB values are missing; reconciliation is incomplete.');continue
                if reported is None or abs(reported-expected)>.02:
                    issue('Review',label+' reconciliation',file,sn,f'{summary_letter}{rr}',detail=f'Reported {reported}; eligible detail {expected:,.2f}. Check section boundaries and exclusions.')
        files.append(dict(file=file,role='PBB form' if count else 'Unrecognized workbook',records=count))
        if not count: issue('Review','No employee records detected',file,detail='Workbook was not audited as a PBB form.')
    byname=defaultdict(list)
    for rec in records: byname[norm(rec['name'])].append(rec)
    for rec in records:
        f,s,r,n=rec['file'],rec['sheet'],rec['row'],rec['name']; cell=f'F{r}'
        matches=byname[norm(n)]
        if len(matches)>1: issue('Review','Duplicate employee',f,s,cell,n,'Also listed at '+ '; '.join(f'{x["file"]} / {x["sheet"]}!F{x["row"]}' for x in matches if x is not rec))
        refs=personnel.get(norm(n),[])
        if not personnel: continue
        if not refs: issue('Review','Personnel name not matched',f,s,cell,n,'No exact normalized full-name match in December 2024 personnel list. Check spelling or personnel status.')
        elif len(refs)>1: issue('Review','Ambiguous personnel match',f,s,cell,n,'Multiple personnel records share this full name.')
        else:
            pf,ps,pr,p=refs[0]
            for label,value,reference in [('Salary grade',rec['sg'],integer(p[6])),('Step',rec['step'],integer(p[7])),('Monthly salary',number(rec['salary']),number(p[5]))]:
                if value is not None and reference is not None and abs(value-reference)>.01: issue('Review','Personnel '+label.lower()+' mismatch',f,s,cell,n,f'Form {value:g}; personnel {reference:g}. Reference: {pf} / {ps}, row {pr}.')
        words=norm(n).split(); first=words[1] if len(words)>1 else ''; last=words[0] if words else ''
        for sf,ss,sr,p in separations:
            sw=norm(p[1]).split()
            if first and sw and sw[0]==first and sw[-1]==last:
                issue('Review','Possible separation match',f,s,cell,n,f'{p[1]} — {p[2]}; last service {p[5]}; effective {p[6]}; remarks {p[7] if len(p)>7 else ""}. Reference: {sf} / {ss}, row {sr}. First-name/surname match only; verify identity and 2024 service. Later-year orders do not establish 2024 ineligibility.')
    if not personnel: issue('Review','Missing personnel reference','',detail='No recognized personnel list; personnel matching was not performed.')
    if not separations: issue('Review','Missing separation reference','',detail='No recognized separation directory; separation matching was not performed.')
    return dict(folder=str(root),files=files,records=len(records),issues=issues,errors=sum(x['level']=='Error' for x in issues),reviews=sum(x['level']=='Review' for x in issues))
