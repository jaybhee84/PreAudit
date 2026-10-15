"""Read PSIPOP incumbent columns and compare PBB employee names locally."""
from pathlib import Path
from collections import defaultdict
from difflib import SequenceMatcher
import re, unicodedata
import csv
from datetime import datetime,date
import pymupdf
import openpyxl

PAY_FOLDER = r'C:\Users\DepEd\Downloads\PBB 2024 WITH IUs'
ROSTER_FOLDER = r'C:\Users\DepEd\Downloads\New folder\09 SEPTEMBER 2026\09 SEPTEMBER 2026'
MEMO_PATH = r"C:\Users\DepEd\Downloads\[MEMORANDUM] DepEd's Eligibility for Fiscal Year 2024 Performance-Based Bonus_r.pdf"
HISTORY_FOLDER = r'C:\Users\DepEd\Desktop\Psipop\Previous Year'
REQUIRED_MONTHS = ['2024-12']+[f'2025-{m:02d}' for m in range(1,7)]
MONTH_NAMES='JANUARY FEBRUARY MARCH APRIL MAY JUNE JULY AUGUST SEPTEMBER OCTOBER NOVEMBER DECEMBER'.split()

def month_from_path(path):
    text=str(path).upper()
    # Prefer the filename / closest folder when multiple dated ancestors exist.
    found=[]
    for i,name in enumerate(MONTH_NAMES,1):
        for match in re.finditer(r'\b'+name+r'\s*[-_ ]*\s*(20\d{2})\b',text):found.append((match.start(),f'{match[1]}-{i:02d}'))
    for match in re.finditer(r'\b(20\d{2})[-_ ](0[1-9]|1[0-2])\b',text):found.append((match.start(),f'{match[1]}-{match[2]}'))
    return max(found)[1] if found else None

def validate_source_folders(roster_folder,history_folder):
    current=Path(roster_folder)
    if current.is_dir():
        months=sorted({month_from_path(path) for path in current.rglob('*.pdf')} - {None})
        if len(months)>1:raise ValueError('Current PSIPOP contains multiple months ('+', '.join(months)+'). Select the Current folder for the latest roster, and put Previous Year in the school-based history field. Historical monthly repetitions are not duplicate employees.')
    if history_folder:
        history=Path(history_folder)
        if history.is_dir():
            months={month_from_path(path) for path in history.rglob('*.pdf')} - {None}
            if months and not months.intersection(REQUIRED_MONTHS):raise ValueError('The school-based history folder contains no December 2024–June 2025 records. The Current and Previous Year folders may be reversed. Put Previous Year in school-based history and Current in the current-roster field.')

def normalized(value):
    s=''.join(c for c in unicodedata.normalize('NFKD',str(value or '')) if not unicodedata.combining(c)).upper()
    return ' '.join(re.findall('[A-Z0-9]+',s))

def candidate_score(pay_name,roster_name):
    """Shared surnames and middle names alone do not establish a name variant."""
    if ',' not in pay_name or ',' not in roster_name:return 0
    ps,pg=pay_name.split(',',1);rs,rg=roster_name.split(',',1)
    ps,rs=normalized(ps),normalized(rs);pg,rg=normalized(pg).split(),normalized(rg).split()
    if not pg or not rg:return 0
    similarity=lambda a,b:SequenceMatcher(None,a,b).ratio()
    first=similarity(pg[0],rg[0])
    if first<.80:return 0
    if ps!=rs and (similarity(ps,rs)<.80 or first<.90):return 0
    # Compare additional given names when both sources contain them; the last token is normally a middle name.
    pcore=pg[:-1] if len(pg)>1 else pg;rcore=rg[:-1] if len(rg)>1 else rg
    for a,b in zip(pcore[1:],rcore[1:]):
        if similarity(a,b)<.80 and not (len(a)==1 and a==b[0]) and not (len(b)==1 and b==a[0]):return 0
    score=similarity(normalized(pay_name),normalized(roster_name))
    return score if score>=.70 else 0

def tin_digits(value):
    if value is None:return ''
    if isinstance(value,float):
        if not value.is_integer():return ''
        value=int(value)
    text=str(value).strip()
    if not re.fullmatch(r'[0-9\s-]+',text):return ''
    digits=re.sub(r'\D','',text)
    return digits if len(digits) in (9,12,13,14) and digits[:9]!='000000000' else ''

def read_tin_reference(file):
    """Optional explicit name-to-TIN source; never infer a TIN from a similar name."""
    if not file:return {},[]
    path=Path(file)
    if not path.is_file():raise ValueError('The name/TIN reference file does not exist.')
    tables=[]
    if path.suffix.lower()=='.csv':
        with path.open(encoding='utf-8-sig',newline='') as stream:
            sample=stream.read(4096);stream.seek(0)
            try:dialect=csv.Sniffer().sniff(sample,delimiters=',;\t')
            except csv.Error:dialect=csv.excel
            tables=[(path.name,list(csv.reader(stream,dialect)))]
    elif path.suffix.lower() in ('.xlsx','.xlsm'):
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        try:tables=[(s.title,list(s.values)) for s in wb]
        finally:wb.close()
    else:raise ValueError('Use an Excel or CSV name/TIN reference with Full Name and TIN columns.')
    result=defaultdict(set);problems=[];found=False
    for title,rows in tables:
        for i,row in enumerate(rows[:50]):
            headers=[normalized(v) for v in row]
            name=next((j for j,v in enumerate(headers) if v in ('NAME','FULL NAME','EMPLOYEE NAME','NAME OF EMPLOYEE','NAME OF INCUMBENT')),None)
            tin=next((j for j,v in enumerate(headers) if v in ('TIN','TAX IDENTIFICATION NUMBER','TAXPAYER IDENTIFICATION NUMBER')),None)
            if name is None or tin is None:continue
            found=True
            for r,values in enumerate(rows[i+1:],i+2):
                if len(values)<=max(name,tin) or not values[name]:continue
                if values[tin] is None or str(values[tin]).strip()=='':continue
                digits=tin_digits(values[tin]);key=normalized(values[name])
                if not digits:problems.append(f'{path.name} / {title}, row {r}: invalid TIN for {values[name]}.');result[key].add('INVALID')
                else:result[key].add(digits[:9])
            break
    if not found:raise ValueError('No Full Name and TIN columns found in the selected TIN reference.')
    return result,problems

def scan_roster(folder,selected_paths=None):
    root=Path(folder)
    if not root.is_dir(): raise ValueError('The PSIPOP folder does not exist.')
    pdfs=sorted(root.rglob('*.pdf')) if selected_paths is None else sorted(selected_paths)
    if not pdfs: raise ValueError('No PSIPOP PDFs found in the selected folder.')
    people=[];coverage=[];problems=[]
    for path in pdfs:
        file=str(path.relative_to(root)); file_people=[];page_checks=[]
        try:
            with pymupdf.open(path) as doc:
                for page_no,page in enumerate(doc,1):
                    page.remove_rotation()
                    spans=[z for b in page.get_text('dict')['blocks'] for line in b.get('lines',[]) for z in line['spans'] if z['text'].strip()]
                    header=next((s for s in spans if 'NAME OF INCUMBENT' in s['text']),None)
                    filled=next((s for s in spans if 'No. of Filled Positions:' in s['text']),None)
                    tin_header=next((s for s in spans if s['text'].strip()=='TIN'),None)
                    if not header or not filled:
                        problems.append(f'{file}, page {page_no}: incumbent column or filled-position subtotal not detected.');continue
                    hy=header['bbox']; left=hy[0]-22;right=hy[2]+22;bottom=filled['bbox'][1]
                    items=sorted([s for s in spans if s['text'].strip().startswith('OSEC-') and s['bbox'][0]<left and hy[3]<s['bbox'][1]<bottom],key=lambda s:s['bbox'][1])
                    found=0
                    for i,item in enumerate(items):
                        y0=item['bbox'][1]-1; y1=items[i+1]['bbox'][1]-1 if i+1<len(items) else bottom-1
                        name_spans=sorted([s for s in spans if left<=s['bbox'][0]<right and y0<=s['bbox'][1]<y1],key=lambda s:(round(s['bbox'][1],1),s['bbox'][0]))
                        name=' '.join(s['text'].strip() for s in name_spans).strip()
                        if not name or normalized(name) in ('VACANT','VACANCY','UNFILLED'):continue
                        if ',' not in name or len(normalized(name).split())<2:
                            problems.append(f'{file}, page {page_no}: could not safely parse incumbent for {item["text"]}.');continue
                        pos=' '.join(s['text'].strip() for s in spans if 175<=s['bbox'][0]<320 and y0<=s['bbox'][1]<y1 and 'space before' not in s['text'])
                        tin_parts=[s['text'].strip() for s in spans if tin_header and tin_header['bbox'][0]-22<=s['bbox'][0]<tin_header['bbox'][2]+20 and y0<=s['bbox'][1]<y1]
                        tin=tin_digits(''.join(tin_parts))
                        file_people.append(dict(name=name,file=file,page=page_no,item=item['text'].strip(),position=pos,tin=tin));found+=1
                    counts=[s for s in spans if filled['bbox'][2]<s['bbox'][0]<240 and abs(s['bbox'][1]-bottom)<3 and re.fullmatch(r'\d+',s['text'].strip())]
                    expected=int(counts[0]['text']) if len(counts)==1 else None
                    page_checks.append(dict(page=page_no,extracted=found,expected=expected))
                    if expected is None or found!=expected:
                        problems.append(f'{file}, page {page_no}: extracted {found} names; printed filled-position subtotal {expected}.')
            people.extend(file_people)
            coverage.append(dict(file=file,pages=len(page_checks),names=len(file_people),checks=page_checks))
        except Exception as exc:problems.append(f'{file}: {exc}');coverage.append(dict(file=file,pages=0,names=0,checks=[]))
    if not people:raise ValueError('No incumbent names could be read from the PSIPOP PDFs.')
    return people,coverage,problems

def read_history(folder):
    if not folder:return dict(people=[],files=[],problems=[],available_months=[],missing_months=REQUIRED_MONTHS,complete_months=[])
    root=Path(folder)
    if not root.is_dir():raise ValueError('The school-based PSIPOP history folder does not exist.')
    paths=[];months={};problems=[]
    for path in root.rglob('*.pdf'):
        month=month_from_path(path)
        if month in REQUIRED_MONTHS:paths.append(path);months[str(path.relative_to(root))]=month
        elif not month:problems.append(f'Historical PDF has no recognizable month/year: {path.relative_to(root)}')
    if not paths:return dict(people=[],files=[],problems=problems+['No December 2024–June 2025 PSIPOP PDFs detected.'],available_months=[],missing_months=REQUIRED_MONTHS,complete_months=[])
    people,files,errors=scan_roster(root,paths);problems.extend(errors)
    for person in people:person.update(month=months[person['file']],source='history')
    for file in files:file['month']=months[file['file']]
    available=sorted(set(months.values()))
    complete=[]
    for month in available:
        group=[f for f in files if f['month']==month]
        if all(f['pages']>0 and all(c['expected']==c['extracted'] for c in f['checks']) for f in group) and not any(any(f['file'] in error for f in group) for error in errors):complete.append(month)
    return dict(people=people,files=files,problems=problems,available_months=available,missing_months=[m for m in REQUIRED_MONTHS if m not in available],complete_months=complete)

def read_payfiles(folder):
    root=Path(folder)
    if not root.is_dir():raise ValueError('The payfile folder does not exist.')
    files=sorted(p for p in root.rglob('*') if p.suffix.lower() in ('.xlsx','.xlsm') and not p.name.startswith('~$'))
    records=[];coverage=[];problems=[]
    for path in files:
        file=str(path.relative_to(root));count=0;recognized=False
        try:
            with_workbook=openpyxl.load_workbook(path,read_only=True,data_only=True)
            try:
                for sheet in with_workbook:
                    rows=list(sheet.values)
                    header=next((i for i,row in enumerate(rows[:45]) if any('NAME OF EMPLOYEE' in str(v).upper() for v in row)),None)
                    if header is None:continue
                    recognized=True
                    col=next(c for c,v in enumerate(rows[header]) if 'NAME OF EMPLOYEE' in str(v).upper())
                    service_col=next((c for c,v in enumerate(rows[header]) if 'MONTHS IN SERVICE' in str(v).upper()),None)
                    salary_col=next((c for c,v in enumerate(rows[header]) if 'MONTHLY BASIC SALARY' in str(v).upper()),None)
                    pbb_col=next((c for c,v in enumerate(rows[header]) if 'AMOUNT OF THE PBB' in str(v).upper()),None)
                    unit=''
                    for r,row in enumerate(rows[header+1:],header+2):
                        if any(re.search(r'ineligible|not eligible',str(v),re.I) for v in row[:5]):break
                        if len(row)>1 and row[1]:unit=str(row[1])
                        name=row[col] if len(row)>col else None
                        has_data=any(v is not None for v in row[col+1:col+7])
                        if not name:
                            # Blank template rows are ignored; personnel data without a name is a coverage failure.
                            if any(isinstance(v,(int,float)) for v in row[col+1:col+3]):problems.append(f'{file} / {sheet.title}, row {r}: employee data has no name.')
                            continue
                        if not has_data or 'TOTAL' in str(name).upper() and ',' not in str(name):continue
                        if ',' not in str(name):
                            problems.append(f'{file} / {sheet.title}, row {r}: name is not in surname, given-name format: {name}.')
                        service=row[service_col] if service_col is not None and len(row)>service_col else None
                        service_text=str(service if service is not None else '')
                        match=re.fullmatch(r'\s*(\d+(?:\.\d+)?)\s*(?:months?)?(?:\s+but\s+less\s+than\s+\d+\s+months?)?\s*',service_text,re.I)
                        months=float(match[1]) if match else None
                        salary=row[salary_col] if salary_col is not None and len(row)>salary_col else None
                        amount=row[pbb_col] if pbb_col is not None and len(row)>pbb_col else None
                        records.append(dict(name=str(name),file=file,sheet=sheet.title,cell=f'{openpyxl.utils.get_column_letter(col+1)}{r}',unit=unit,service_label=service_text,service_months=months,service_cell=f'{openpyxl.utils.get_column_letter(service_col+1)}{r}' if service_col is not None else '',salary=salary if isinstance(salary,(int,float)) and not isinstance(salary,bool) else None,payfile_pbb=amount if isinstance(amount,(int,float)) and not isinstance(amount,bool) else None));count+=1
            finally:with_workbook.close()
            coverage.append(dict(file=file,records=count,role='Payfile' if recognized else 'Reference / not a payfile'))
        except Exception as exc:problems.append(f'{file}: {exc}');coverage.append(dict(file=file,records=0,role='Unreadable'))
    if not records:raise ValueError('No employee entries detected. Select the folder containing the PBB Form 1.0 payfiles.')
    return records,coverage,problems

def retirement_date(value):
    if isinstance(value,datetime):return value.date().isoformat()
    if isinstance(value,date):return value.isoformat()
    text=re.sub(r'\s+',' ',str(value or '').replace(',',' ').strip()).upper()
    for pattern in ('%B %d %Y','%b %d %Y','%Y-%m-%d','%m/%d/%Y'):
        try:return datetime.strptime(text,pattern).date().isoformat()
        except ValueError:pass
    return None

def read_retirements(folder):
    records=[]
    for path in Path(folder).rglob('*.xlsx'):
        if 'SEPARATION' not in path.name.upper() or path.name.startswith('~$'):continue
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        try:tables={s.title:list(s.values) for s in wb}
        finally:wb.close()
        item_by_name=defaultdict(set)
        for title,rows in tables.items():
            if 'DIRECTORY' in title.upper():continue
            if len(rows)>22 and len(rows[19])>5 and len(rows[22])>5 and rows[19][5] and rows[22][5]:item_by_name[normalized(rows[19][5])].add(str(rows[22][5]))
        for title,rows in tables.items():
            if 'DIRECTORY' not in title.upper():continue
            for i,row in enumerate(rows,1):
                if i<=12 or len(row)<7 or 'RETIRE' not in str(row[2]).upper():continue
                if len(row)>7 and 'CANCEL' in str(row[7]).upper():continue
                last=retirement_date(row[5]);effective=retirement_date(row[6]);cutoff=last or effective
                if cutoff:records.append(dict(name=str(row[1]),last_service=last,effective=effective,date=cutoff,items=sorted(item_by_name.get(normalized(row[1]),set())),file=str(path.relative_to(folder)),sheet=title,cell=f'B{i}'))
    return records

def compare(pay_folder=PAY_FOLDER,roster_folder=ROSTER_FOLDER,tin_file=None,history_folder=None):
    validate_source_folders(roster_folder,history_folder)
    roster,roster_files,roster_problems=scan_roster(roster_folder)
    pay,pay_files,pay_problems=read_payfiles(pay_folder)
    tin_reference,tin_problems=read_tin_reference(tin_file)
    problems=roster_problems+pay_problems+tin_problems
    tins=defaultdict(list)
    for person in roster:
        if person.get('tin'):tins[person['tin'][:9]].append(person)
    index=defaultdict(list)
    for person in roster:index[normalized(person['name'])].append(person)
    duplicates=defaultdict(int)
    for rec in pay:duplicates[normalized(rec['name'])]+=1
    results=[]
    for rec in pay:
        key=normalized(rec['name']);exact=index.get(key,[]);candidates=[]
        if exact:
            status='Matched' if len(exact)==1 else 'Multiple roster entries'
            refs=exact
        else:
            for k,people in index.items():
                score=candidate_score(rec['name'],people[0]['name'])
                if score:candidates.append((score,people))
            candidates.sort(key=lambda x:x[0],reverse=True)
            refs=[p for _,group in candidates[:3] for p in group]
            status='Possible name variant' if refs else ('Name format review' if ',' not in rec['name'] else 'Check roster coverage' if roster_problems else 'Separated / absent from current PSIPOP')
        detail=('Exact normalized full-name match.' if exact else 'Similar names require confirmation; not counted as matched.' if refs else 'Name is not in surname, given-name format; verify before classifying.' if status=='Name format review' else 'No current-roster match. Roster extraction has unresolved coverage issues.' if roster_problems else 'Absent from September 2026 PSIPOP; classified as separated under your current-roster rule. This alone does not determine FY 2024 PBB eligibility.')
        pay_tins=tin_reference.get(key,set());pay_tin='';basis='Name'
        if pay_tins:
            basis='TIN reference'
            if len(pay_tins)!=1 or 'INVALID' in pay_tins:status='TIN needs review';detail='Name/TIN reference contains invalid or conflicting TINs for this payfile name.'
            else:
                pay_tin=next(iter(pay_tins));tin_matches=tins.get(pay_tin,[])
                if len(tin_matches)>1:status='Multiple TIN matches';refs=tin_matches;detail='The TIN appears on more than one current PSIPOP entry; verify identity.'
                elif len(tin_matches)==1:
                    target=tin_matches[0]
                    if exact and all(normalized(p['name'])!=normalized(target['name']) for p in exact):status='TIN/name conflict';refs=exact+[target];detail='TIN points to a different current person than the exact name. Check the source TIN.'
                    else:status='TIN matched';refs=tin_matches;detail='Unique current PSIPOP match using the 9-digit taxpayer identifier; a branch suffix is excluded.'
                else:status='TIN not found';detail='The source TIN was not found in the current PSIPOP. Verify the TIN and personnel identity before classifying.'
        results.append(dict(**rec,status=status,duplicate=duplicates[key]>1,references=refs,detail=detail,tin=pay_tin,match_basis=basis))
    counts={s:sum(r['status']==s for r in results) for s in sorted(set(r['status'] for r in results))}
    history=read_history(history_folder)
    return dict(results=results,roster=roster,history=history,retirements=read_retirements(pay_folder),required_months=REQUIRED_MONTHS,history_folder=history_folder or '',counts=counts,payfiles=pay_files,rosterfiles=roster_files,roster_names=len(roster),roster_unique=len(index),pay_entries=len(pay),problems=problems,roster_problems=roster_problems,pay_folder=pay_folder,roster_folder=roster_folder,tin_file=tin_file or '',roster_tins=sum(bool(p.get('tin')) for p in roster))

def export_workbook(file_path,sheets):
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table,TableStyleInfo
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    workbook=openpyxl.Workbook()
    workbook.remove(workbook.active)
    for number,sheet in enumerate(sheets,1):
        ws=workbook.create_sheet(sheet['name'])
        for values in sheet['rows']:
            ws.append([ILLEGAL_CHARACTERS_RE.sub('',value) if isinstance(value,str) else value for value in values])
            for cell in ws[ws.max_row]:
                # Employee names and notes remain literal text, including leading '='.
                if isinstance(cell.value,str):cell.data_type='s'
        ws.freeze_panes='E2'
        ws.sheet_view.showGridLines=False
        ws.row_dimensions[1].height=32
        for cell in ws[1]:
            cell.font=Font(name='Calibri',bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='143C43')
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        for col in range(1,ws.max_column+1):
            title=str(ws.cell(1,col).value or '')
            width=35 if title in ['Payfile employee','Delivery unit','Remarks'] else 27 if title in ['Payfile','PSIPOP references','Selected PSIPOP names'] else 20
            ws.column_dimensions[get_column_letter(col)].width=width
        for cells in ws.iter_rows(min_row=2):
            for cell in cells:cell.font=Font(name='Calibri',size=11);cell.alignment=Alignment(vertical='top')
            ws.row_dimensions[cells[0].row].height=18
        if ws.max_row>1:
            table=Table(displayName=f'Personnel{number}',ref=ws.dimensions)
            table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            ws.add_table(table)
        else:ws.auto_filter.ref=ws.dimensions
        titles=[cell.value for cell in ws[1]]
        for title,fmt in [('Service factor','0%'),('Indicative PBB (payfile salary basis)','#,##0.00')]:
            if title in titles:
                col=titles.index(title)+1
                for row in range(2,ws.max_row+1):ws.cell(row,col).number_format=fmt
    workbook.save(file_path)
    return {'file':str(file_path),'sheets':{ws.title:max(ws.max_row-1,0) for ws in workbook.worksheets}}

if __name__=='__main__':
    import json,sys
    try:
        args=json.load(sys.stdin)
        result=export_workbook(args['file'],args['sheets']) if args.get('mode')=='export_workbook' else compare(args.get('pay_folder',PAY_FOLDER),args.get('roster_folder',ROSTER_FOLDER),args.get('tin_file'),args.get('history_folder'))
        print(json.dumps(result,ensure_ascii=True))
    except Exception as exc:
        print(json.dumps({'error':str(exc)}));sys.exit(1)
