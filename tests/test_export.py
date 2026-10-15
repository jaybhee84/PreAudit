import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
import openpyxl
from roster import export_workbook

class WorkbookExportTests(unittest.TestCase):
    def test_tabs_literal_text_numbers_and_filters(self):
        with TemporaryDirectory() as folder:
            path=Path(folder)/'export.xlsx'
            result=export_workbook(path,[{'name':'IUs','rows':[['Name','Service factor','Indicative PBB (payfile salary basis)','Remarks'],['=1+1',0.8,11700,'Prorated 80%']]},{'name':'SDO','rows':[['Name']]},{'name':'School-based','rows':[['Name'],['Teacher']]}])
            book=openpyxl.load_workbook(path)
            self.assertEqual(book.sheetnames,['IUs','SDO','School-based'])
            self.assertEqual(book['IUs']['A2'].data_type,'s')
            self.assertEqual(book['IUs']['A2'].value,'=1+1')
            self.assertEqual(book['IUs']['B2'].value,.8)
            self.assertEqual(book['IUs']['B2'].number_format,'0%')
            self.assertEqual(book['IUs'].freeze_panes,'E2')
            self.assertEqual(len(book['IUs'].tables),1)
            self.assertEqual(result['sheets'],{'IUs':1,'SDO':0,'School-based':1})
            book.close()
